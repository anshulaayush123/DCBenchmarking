"""Create a new version of the DC benchmarking Excel workbook after an approved update.

Runs on the owner's laptop (Windows Task Scheduler starts it once a day). It downloads
data.json from the `main` branch of the GitHub repo, which only changes when the owner
merges a pull request. If data.json differs from the copy saved at the last run, it:

  1. opens the latest workbook in the output folder,
  2. writes every changed value into the "Benchmarking" tab and highlights it with a
     note "Updated DD/MM/YYYY (was ...)",
  3. rebuilds the "Sources" tab: the same grid, where each cell shows how the value was
     calculated and links to the source document,
  4. saves the result as DC_benchmarking_Claude_DDMMYYYY_vX.xlsx in the same folder.

Older files are never changed. Cells that the dashboard does not track (for example
MasTec or the HR and Corporate Functions headcount rows) are left as they are.

Usage:
  python update_excel.py                       normal run (what the scheduled task does)
  python update_excel.py --old A.json --new B.json --workbook X.xlsx --output-dir DIR
                                               compare two local files, for testing
"""

import argparse
import copy
import datetime as dt
import json
import logging
import os
import re
import sys
import urllib.parse
import urllib.request
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

DATA_URL = "https://raw.githubusercontent.com/anshulaayush123/DCBenchmarking/main/data.json"
OUTPUT_DIR = (r"C:\Users\AnshulApurva_\OneDrive - Data Volt Investment LLC\Desktop"
              r"\Anshul + Usamah\Internal Strategy materials\DCs Bechmarking exercise"
              r"\Claude versions")
STATE_DIR = Path(os.environ.get("LOCALAPPDATA", Path.home())) / "DCBenchmarking"

DATA_SHEET = "Benchmarking"
SOURCES_SHEET = "Sources"
HEADER_ROW = 3          # company names
LABEL_COL = 2           # metric names (column B)
FIRST_DATA_COL = 3      # column C

HIGHLIGHT = PatternFill("solid", fgColor="FFFFEB9C")
WHITE = "FFFFFFFF"
DARK = "FF1F1F1F"
LINK_FONT = Font(color="FF0563C1", underline="single", size=9)
NOTE_FONT = Font(color="FF595959", size=9)
FILE_RE = re.compile(r"^DC_benchmarking_Claude_(\d{2})(\d{2})(\d{4})_v(\d+)\.xlsx$", re.I)

# data.json metric names whose Excel row label is worded differently.
ALIASES = {
    "netdebtebitda": "netdebttoebitdax",
    "netdebtequity": "netdebttoequityx",
    "patmw": "patmmw",
    "netdebtmw": "netdebtmmw",
    "totalemployees": "totalofemployees",
}

# Derived metrics, as defined in AGENT_INSTRUCTIONS.md.
DERIVED = {
    "Live vs pipeline %": "Installed MW ÷ Pipeline MW",
    "EBITDA Margin": "EBITDA ÷ Revenue",
    "PAT Margin": "PAT ÷ Revenue",
    "Net Debt/EBITDA": "Net Debt ÷ EBITDA (N/A if EBITDA ≤ 0)",
    "G&A % of Revenue": "G&A ÷ Revenue",
    "EBITDA/MW": "EBITDA ÷ installed MW",
    "PAT/MW": "PAT ÷ installed MW",
    "Net debt/MW": "Net Debt ÷ installed MW",
    "MW/FTE": "Installed MW ÷ total employees",
    "Revenue ($M)/FTE": "Revenue ÷ total employees",
    "EBITDA ($M)/FTE": "EBITDA ÷ total employees",
}

log = logging.getLogger("update_excel")


def norm(label):
    key = re.sub(r"[^a-z0-9]", "", str(label).lower())
    return ALIASES.get(key, key)


def flatten(data):
    """{(company, metric): value} for every metric, plus each company's period."""
    values = {(c["name"], "Year"): c.get("period") for c in data.get("companies", [])}
    for metrics in data.get("metrics", {}).values():
        for metric, per_company in metrics.items():
            for company, value in per_company.items():
                values[(company, metric)] = value
    return values


def to_cell_value(metric, value):
    """Convert a data.json display value into what the Excel cell should hold."""
    if isinstance(value, (int, float)):
        # Headcount splits are whole percentages in data.json, fractions in Excel.
        return value / 100 if metric.endswith(" %") and metric != "Live vs pipeline %" else value
    if not isinstance(value, str):
        return value
    text = value.strip()
    number = text.replace(",", "")
    if re.fullmatch(r"-?\d+(\.\d+)?%", number):
        return float(number[:-1]) / 100
    if re.fullmatch(r"-?\d+(\.\d+)?x", number):
        return float(number[:-1])
    if re.fullmatch(r"-?\d+(\.\d+)?", number):
        return float(number)
    return text


def describe(value, number_format=""):
    """Show an old cell value the way the cell displays it."""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if "%" in number_format:
            return f"{value:.1%}"
        if "x" in number_format:
            return f"{value:.1f}x"
        return f"{value:,.2f}".rstrip("0").rstrip(".") if abs(value) < 100 else f"{value:,.0f}"
    return "blank" if value in (None, "") else str(value)


def short_link(url):
    """'https://www.sec.gov/Archives/.../eqix-20260630.htm' -> 'sec.gov: eqix-20260630.htm'."""
    parts = urllib.parse.urlsplit(url)
    host = parts.netloc.removeprefix("www.")
    name = urllib.parse.unquote(parts.path.rstrip("/").rsplit("/", 1)[-1])
    return f"{host}: {name}" if name else host


def grid(ws):
    """Map company -> column and normalized metric label -> row on the data sheet."""
    columns = {}
    for col in range(FIRST_DATA_COL, ws.max_column + 1):
        name = ws.cell(HEADER_ROW, col).value
        if name:
            columns[str(name).strip()] = col
    rows = {}
    for row in range(HEADER_ROW + 1, ws.max_row + 1):
        label = ws.cell(row, LABEL_COL).value
        if label:
            rows.setdefault(norm(label), row)
    return columns, rows


def find_base_workbook(folder):
    """Latest DC_benchmarking_Claude_* file, else the newest other .xlsx in the folder."""
    claude, others = [], []
    for path in Path(folder).glob("*.xlsx"):
        if path.name.startswith("~$"):
            continue
        m = FILE_RE.match(path.name)
        if m:
            day, month, year, version = map(int, m.groups())
            claude.append(((year, month, day, version), path))
        else:
            others.append((path.stat().st_mtime, path))
    if claude:
        return max(claude)[1]
    if others:
        return max(others)[1]
    return None


def next_output_path(folder, today):
    stamp = today.strftime("%d%m%Y")
    taken = [int(m.group(4)) for p in Path(folder).glob("*.xlsx")
             if (m := FILE_RE.match(p.name)) and "".join(m.groups()[:3]) == stamp]
    return Path(folder) / f"DC_benchmarking_Claude_{stamp}_v{max(taken, default=0) + 1}.xlsx"


def clear_old_highlights(ws):
    """Give last update's highlighted cells back the normal fill of their row."""
    for row in ws.iter_rows(min_row=HEADER_ROW + 1, min_col=FIRST_DATA_COL):
        for cell in row:
            if cell.fill.fgColor.rgb == HIGHLIGHT.fgColor.rgb:
                label = ws.cell(cell.row, LABEL_COL)
                cell.fill = copy.copy(label.fill)
                if cell.font.color is not None and cell.font.color.rgb == DARK:
                    cell.font = copy.copy(label.font)


def apply_changes(ws, changes, today):
    columns, rows = grid(ws)
    written, skipped = [], []
    for (company, metric), (old, new) in sorted(changes.items()):
        col, row = columns.get(company), rows.get(norm(metric))
        if not col or not row:
            skipped.append(f"{company} / {metric}")
            continue
        cell = ws.cell(row, col)
        was = cell.value
        cell.value = to_cell_value(metric, new)
        cell.fill = HIGHLIGHT
        if cell.font.color is not None and cell.font.color.rgb == WHITE:
            # White header text (the period row) is unreadable on yellow.
            font = copy.copy(cell.font)
            font.color = DARK
            cell.font = font
        cell.comment = Comment(
            f"Updated {today:%d/%m/%Y} (was {describe(was, cell.number_format)})", "Claude",
            width=220, height=60)
        written.append((row, col))
    return written, skipped


def build_sources(wb, data, changed_cells, today):
    """Rebuild the Sources tab: same layout as the data tab, one note per value."""
    data_ws = wb[DATA_SHEET]
    if SOURCES_SHEET in wb.sheetnames:
        del wb[SOURCES_SHEET]
    ws = wb.create_sheet(SOURCES_SHEET, index=wb.sheetnames.index(DATA_SHEET) + 1)

    sources = data.get("sources", {})
    methods = data.get("methods", {})
    tracked = {norm(m): m for metrics in data.get("metrics", {}).values() for m in metrics}
    tracked[norm("Year")] = "Year"
    companies = {c["name"] for c in data.get("companies", [])}

    # Copy titles, headers and the metric labels with their styling.
    for row in data_ws.iter_rows(max_col=LABEL_COL):
        for src in row:
            dst = ws.cell(src.row, src.column, src.value)
            dst._style = copy.copy(src._style)
    for src in data_ws[HEADER_ROW]:
        dst = ws.cell(HEADER_ROW, src.column, src.value)
        dst._style = copy.copy(src._style)
    for merged in data_ws.merged_cells.ranges:
        ws.merge_cells(str(merged))
    ws["A1"] = "Data Center Operators Benchmarking: sources and calculation method"
    ws["A2"] = ("Each cell explains the matching cell on the Benchmarking tab. "
                "Click a blue cell to open its source document.")
    ws.column_dimensions["A"].width = data_ws.column_dimensions["A"].width or 20
    ws.column_dimensions["B"].width = data_ws.column_dimensions["B"].width or 35
    ws.freeze_panes = "C4"
    for row in range(HEADER_ROW + 1, data_ws.max_row + 1):
        ws.row_dimensions[row].height = 80

    for col in range(FIRST_DATA_COL, data_ws.max_column + 1):
        company = data_ws.cell(HEADER_ROW, col).value
        if not company:
            continue
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 45
        for row in range(HEADER_ROW + 1, data_ws.max_row + 1):
            label = data_ws.cell(row, LABEL_COL).value
            if not label:
                continue
            metric = tracked.get(norm(label))
            url = sources.get(company, {}).get(metric) if metric else None
            method = methods.get(company, {}).get(metric) if metric else None
            if company not in companies or metric is None:
                text = "Not covered by the update agent; value from the original workbook."
            elif metric == "Year":
                text = "Reporting period of this company's figures."
            elif method or url:
                text = method or "Taken from the source document."
            elif metric in DERIVED:
                text = f"Calculated: {DERIVED[metric]}"
            else:
                text = "Original workbook value; source not recorded."
            if (row, col) in changed_cells:
                text = f"Updated {today:%d/%m/%Y}. {text}"
            cell = ws.cell(row, col)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = NOTE_FONT
            if url:
                cell.value = f"{text}\nSource: {short_link(url)} (click to open)"
                cell.hyperlink = url
                cell.font = LINK_FONT
            else:
                cell.value = text
            if (row, col) in changed_cells:
                cell.fill = HIGHLIGHT


def fetch_remote():
    request = urllib.request.Request(DATA_URL, headers={"Cache-Control": "no-cache"})
    with urllib.request.urlopen(request, timeout=60) as response:
        return json.loads(response.read().decode("utf-8"))


def run(old, new, base, output_dir, today):
    old_values, new_values = flatten(old), flatten(new)
    changes = {key: (old_values.get(key), value) for key, value in new_values.items()
               if old_values.get(key) != value}
    if not changes:
        log.info("data.json has not changed since the last run; nothing to do.")
        return None

    wb = openpyxl.load_workbook(base)
    ws = wb[DATA_SHEET]
    clear_old_highlights(ws)
    written, skipped = apply_changes(ws, changes, today)
    for item in skipped:
        log.warning("No matching cell in the workbook for %s; skipped.", item)
    build_sources(wb, new, set(written), today)
    for name in wb.sheetnames:
        wb[name].sheet_view.tabSelected = name == DATA_SHEET
    wb.active = wb.sheetnames.index(DATA_SHEET)

    out = next_output_path(output_dir, today)
    wb.save(out)
    log.info("Saved %s (%d values changed, based on %s).", out.name, len(written), base.name)
    return out


def main():
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--old", type=Path, help="previous data.json (default: saved copy)")
    parser.add_argument("--new", type=Path, help="new data.json (default: download from main)")
    parser.add_argument("--workbook", type=Path, help="workbook to start from")
    parser.add_argument("--output-dir", type=Path, default=Path(OUTPUT_DIR))
    args = parser.parse_args()

    STATE_DIR.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(
        level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s",
        handlers=[logging.FileHandler(STATE_DIR / "update_excel.log", encoding="utf-8"),
                  logging.StreamHandler(sys.stdout)])
    state_file = STATE_DIR / "last_data.json"
    testing = args.old is not None or args.new is not None

    try:
        new = json.loads(args.new.read_text("utf-8")) if args.new else fetch_remote()
        if args.old:
            old = json.loads(args.old.read_text("utf-8"))
        elif state_file.exists():
            old = json.loads(state_file.read_text("utf-8"))
        else:
            state_file.write_text(json.dumps(new, indent=2), "utf-8")
            log.info("First run: saved today's data.json as the starting point. "
                     "The next approved update will create a new workbook.")
            return 0

        if not args.output_dir.is_dir():
            log.error("Output folder not found: %s", args.output_dir)
            return 1
        base = args.workbook or find_base_workbook(args.output_dir)
        if base is None:
            log.error("No workbook found in %s. Put the latest Excel file there.",
                      args.output_dir)
            return 1

        run(old, new, base, args.output_dir, dt.date.today())
        if not testing:
            state_file.write_text(json.dumps(new, indent=2), "utf-8")
        return 0
    except Exception:
        log.exception("Update failed; data.json state left unchanged so it retries next run.")
        return 1


if __name__ == "__main__":
    sys.exit(main())
